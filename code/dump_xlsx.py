import glob, os, csv, wave
from openpyxl import load_workbook

D = os.path.expanduser("~/vsip/data/bubble_wave_1031/bubble wave-2023.10.31")
outdir = os.path.expanduser("~/vsip/results/ref_errors")
os.makedirs(outdir, exist_ok=True)

# 1) wav header info (Fs sanity check)
p = sorted(glob.glob(os.path.join(D, "*.wav")))[0]
with wave.open(p, "rb") as w:
    print("first wav:", os.path.basename(p), "ch=", w.getnchannels(),
          "fs=", w.getframerate(), "frames=", w.getnframes(), "width=", w.getsampwidth())

# 2) xlsx -> csv
for p in sorted(glob.glob(os.path.join(D, "*.xlsx"))):
    wb = load_workbook(p, data_only=True)
    base = os.path.splitext(os.path.basename(p))[0].replace(" ", "_")
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        while rows and all(v is None for v in rows[-1]):
            rows.pop()
        out = os.path.join(outdir, f"{base}__{ws.title.replace('/', '_')}.csv")
        with open(out, "w", newline="") as f:
            cw = csv.writer(f)
            for r in rows:
                cw.writerow(["" if v is None else v for v in r])
        print(out, len(rows), "rows x", max(len(r) for r in rows) if rows else 0, "cols")
